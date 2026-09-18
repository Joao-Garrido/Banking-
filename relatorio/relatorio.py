"""Constrói o livro de consolidação da família — folhas, tabelas e gráficos.

    python relatorio/relatorio.py --exemplo
    python relatorio/relatorio.py /caminho/para/fonte --saida out/consolidado.xlsx

O livro sai organizado como um dossier, do resumo para a prova:

    Capa · Alocação · Posições · Movimentos · Notas

Este módulo é só montagem: não escolhe cores, não escolhe formatos, não decide
espaçamentos — pede tudo isso a `design.py`. Acrescentar uma folha é escrever
aqui uma função `folha_*`, consumir os tokens e registá-la em `construir()`.
"""

from __future__ import annotations

import argparse
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

try:  # corre como script (python relatorio/relatorio.py) ou como módulo
    from . import dados, design
except ImportError:  # pragma: no cover
    import dados
    import design

__all__ = ["construir"]


def _f(valor: Decimal | None) -> float | None:
    """Decimal → float, só na fronteira com o Excel (o Excel não tem Decimal)."""
    return None if valor is None else float(valor)


# ── Capa ──────────────────────────────────────────────────────────────────────
def folha_capa(wb: Workbook, rel: dados.Relatorio, arte: Path) -> None:
    ws = wb.create_sheet("Capa")
    design.preparar_folha(ws, larguras=[2.5] + [11] * 16, tab=design.INK)
    periodo = f"{rel.inicio:%d/%m/%Y} a {rel.fim:%d/%m/%Y} · valores em {rel.moeda}"
    linha = design.titulo_folha(ws, f"{rel.familia} — consolidação", periodo)

    variacao = None
    if len(rel.serie) >= 2 and rel.serie[0].patrimonio:
        variacao = float((rel.serie[-1].patrimonio / rel.serie[0].patrimonio - 1) * 100)
    liquidez = sum((p.valor for p in rel.posicoes if p.classe == "Liquidez"), Decimal("0"))
    entradas = sum((m.valor for m in rel.movimentos if m.valor > 0), Decimal("0"))

    linha = design.faixa_kpi(ws, linha, 2, [
        dict(rotulo=f"Património ({rel.moeda})", valor=_f(rel.patrimonio),
             delta=variacao, delta_rotulo="que no início do período"),
        dict(rotulo="Contas consolidadas", valor=len(rel.contas),
             formato="0", delta_rotulo=f"{len(rel.posicoes)} posições"),
        dict(rotulo="Liquidez", valor=_f(liquidez),
             delta_rotulo=f"{float(rel.peso(liquidez)):.1f}% da carteira"),
        dict(rotulo="Entradas do período", valor=_f(entradas),
             delta_rotulo=f"vs. {rel.referencia}" if rel.referencia else ""),
    ])

    classes = rel.por_classe()
    design.colar_grafico(ws, design.grafico_linha(
        arte / "evolucao.png",
        [f"{s.mes:%b}" for s in rel.serie],
        [("Património", [_f(s.patrimonio) for s in rel.serie])],
        titulo="Evolução do património",
        subtitulo=f"fecho de cada mês, {rel.moeda}",
    ), f"B{linha + 1}", 7.6, 3.6)

    design.colar_grafico(ws, design.grafico_rosca(
        arte / "alocacao.png",
        [c for c, _ in classes], [_f(v) for _, v in classes],
        titulo="Alocação por classe",
        subtitulo="peso sobre o património consolidado",
        centro=f"{float(rel.patrimonio)/1_000_000:.1f}M",
        centro_rotulo=rel.moeda,
    ), f"K{linha + 1}", 5.2, 4.0)

    linha += 21  # altura do gráfico mais alto da fila (4,0 pol ≈ 19 linhas) + folga
    design.colar_grafico(ws, design.grafico_cascata(
        arte / "cascata.png",
        [(a.rubrica, _f(a.valor), a.e_total) for a in rel.atribuicao],
        titulo="De onde veio a variação do período",
        subtitulo="abertura, fluxos e efeito de mercado até ao fecho",
    ), f"B{linha}", 7.6, 3.8)

    maiores = rel.maiores_posicoes(8)
    design.colar_grafico(ws, design.grafico_ranking(
        arte / "maiores.png",
        [p.titulo for p in maiores], [_f(p.valor) for p in maiores],
        titulo="Maiores posições",
        subtitulo=f"valor de mercado, {rel.moeda}",
    ), f"K{linha}", 6.4, 4.0)


# ── Alocação ──────────────────────────────────────────────────────────────────
def folha_alocacao(wb: Workbook, rel: dados.Relatorio, arte: Path) -> None:
    ws = wb.create_sheet("Alocação")
    design.preparar_folha(ws, larguras=[2.5, 24, 16, 12, 16] + [11] * 12, tab=design.SERIES[0])
    linha = design.titulo_folha(ws, "Alocação", f"por classe e por conta · {rel.moeda}")

    classes = rel.por_classe()
    corpo = [[classe, _f(valor), float(rel.peso(valor)),
              len([p for p in rel.posicoes if p.classe == classe])]
             for classe, valor in classes]
    fim = design.escrever_tabela(
        ws, linha, 2,
        ["Classe", f"Valor ({rel.moeda})", "Peso", "Posições"],
        corpo,
        estilos=["fo_corpo", "fo_dinheiro", "fo_percent", "fo_corpo"],
        alinhamentos=["left", "right", "right", "right"],
        total=["Total", _f(rel.patrimonio), 100.0, len(rel.posicoes)],
    )
    design.barra_de_peso(ws, f"D{linha + 1}:D{linha + len(corpo)}")

    contas, series = rel.por_conta_e_classe()
    design.colar_grafico(ws, design.grafico_barras_empilhadas(
        arte / "alocacao_contas.png",
        contas, [(nome, [_f(v) for v in valores]) for nome, valores in series],
        titulo="Composição de cada conta",
        subtitulo=f"valor de mercado, {rel.moeda}",
        unidade="",
    ), f"G{linha}", 7.6, 4.0)

    nota = ws.cell(row=fim + 1, column=2,
                   value="A cor de cada classe é fixa: a mesma classe tem o mesmo tom "
                         "em todas as folhas e em todos os gráficos.")
    nota.font = design.Font(name=design.FONT, size=8, color=design.xl(design.MUTED))


# ── Posições ──────────────────────────────────────────────────────────────────
def folha_posicoes(wb: Workbook, rel: dados.Relatorio) -> None:
    ws = wb.create_sheet("Posições")
    design.preparar_folha(ws, larguras=[2.5, 12, 16, 34, 14, 12, 14, 16, 16, 10], tab=design.SERIES[2])
    linha = design.titulo_folha(ws, "Posições", f"todas as contas, {rel.fim:%d/%m/%Y}")

    corpo = [
        [p.conta, p.classe, p.titulo, _f(p.quantidade), _f(p.preco),
         _f(p.custo), _f(p.valor), _f(p.mais_valia), float(rel.peso(p.valor))]
        for p in sorted(rel.posicoes, key=lambda p: p.valor, reverse=True)
    ]
    fim = design.escrever_tabela(
        ws, linha, 2,
        ["Conta", "Classe", "Título", "Quantidade", "Preço", "Custo",
         f"Valor ({rel.moeda})", "Mais-valia", "Peso"],
        corpo,
        estilos=["fo_corpo", "fo_corpo", "fo_corpo", "fo_corpo", "fo_dinheiro",
                 "fo_dinheiro", "fo_dinheiro", "fo_dinheiro", "fo_percent"],
        alinhamentos=["left", "left", "left", "right", "right", "right",
                      "right", "right", "right"],
        total=["Total", None, None, None, None, None, _f(rel.patrimonio), None, 100.0],
    )
    design.barra_de_peso(ws, f"J{linha + 1}:J{linha + len(corpo)}", cor=design.SERIES[2])
    ws.cell(row=fim + 1, column=2,
            value="Campo vazio = a fonte não o imprime. Nada aqui é estimado."
            ).font = design.Font(name=design.FONT, size=8, color=design.xl(design.MUTED))


# ── Movimentos ────────────────────────────────────────────────────────────────
def folha_movimentos(wb: Workbook, rel: dados.Relatorio) -> None:
    ws = wb.create_sheet("Movimentos")
    design.preparar_folha(ws, larguras=[2.5, 14, 12, 18, 40, 16], tab=design.SERIES[1])
    linha = design.titulo_folha(ws, "Movimentos", f"{rel.inicio:%d/%m/%Y} a {rel.fim:%d/%m/%Y}")

    corpo = [[m.data, m.conta, m.tipo, m.descricao, _f(m.valor)]
             for m in sorted(rel.movimentos, key=lambda m: m.data)]
    design.escrever_tabela(
        ws, linha, 2,
        ["Data", "Conta", "Tipo", "Descrição", f"Valor ({rel.moeda})"],
        corpo,
        estilos=["fo_corpo", "fo_corpo", "fo_corpo", "fo_corpo", "fo_dinheiro"],
        alinhamentos=["left", "left", "left", "left", "right"],
        total=["Total", None, None, None, _f(sum((m.valor for m in rel.movimentos), Decimal("0")))],
    )
    for r in range(linha + 1, linha + 1 + len(corpo)):
        ws.cell(row=r, column=2).number_format = design.FMT_DATA


# ── Notas ─────────────────────────────────────────────────────────────────────
def folha_notas(wb: Workbook, rel: dados.Relatorio, problemas: list[str]) -> None:
    ws = wb.create_sheet("Notas")
    design.preparar_folha(ws, larguras=[2.5, 90], tab=design.MUTED, paisagem=False)
    linha = design.titulo_folha(ws, "Notas e pressupostos",
                           "o que sustenta os números das folhas anteriores")

    def paragrafo(texto: str, cor: str = design.INK_2, tamanho: int = 10, negrito: bool = False):
        nonlocal linha
        c = ws.cell(row=linha, column=2, value=texto)
        c.font = design.Font(name=design.FONT, size=tamanho, bold=negrito, color=design.xl(cor))
        c.alignment = design.Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[linha].height = 18 + 12 * (len(texto) // 95)
        linha += 1

    paragrafo("Consistência", design.INK, 12, True)
    if problemas:
        for p in problemas:
            paragrafo(f"{design.ICONE_ESTADO['critico']} contradição — {p}", design.ESTADO["critico"])
    else:
        paragrafo(f"{design.ICONE_ESTADO['bom']} sem contradições — a série fecha com as posições "
                  f"e a cascata fecha com a abertura.", design.DELTA_BOM)
    linha += 1

    paragrafo("Fonte e método", design.INK, 12, True)
    for nota in rel.notas:
        paragrafo(f"· {nota}")
    paragrafo(f"· Contas consolidadas: "
              f"{'; '.join(f'{c.id} ({c.custodiante}, {c.moeda})' for c in rel.contas)}.")
    paragrafo(f"· Moeda de reporte: {rel.moeda}. Posições em moeda estrangeira convertidas "
              f"à taxa da data de posição.")
    paragrafo("· Campo vazio significa que a fonte não o imprime — não significa zero.")


# ── Montagem ──────────────────────────────────────────────────────────────────
def construir(rel: dados.Relatorio, destino: Path) -> Path:
    """Escreve o livro e os PNG dos gráficos ao lado. Devolve o caminho do livro."""
    destino = Path(destino)
    arte = destino.parent / "graficos"
    problemas = dados.validar(rel)

    wb = Workbook()
    wb.remove(wb.active)
    design.registar_estilos(wb)

    folha_capa(wb, rel, arte)
    folha_alocacao(wb, rel, arte)
    folha_posicoes(wb, rel)
    folha_movimentos(wb, rel)
    folha_notas(wb, rel, problemas)

    wb.properties.title = f"{rel.familia} — consolidação {rel.fim:%Y}"
    wb.properties.creator = "relatorio/"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def main() -> None:
    p = argparse.ArgumentParser(description="Relatório de consolidação de family office")
    p.add_argument("origem", nargs="?", type=Path,
                   help="fonte de dados (ver dados.carregar)")
    p.add_argument("--exemplo", action="store_true",
                   help="usa a carteira sintética em vez de uma fonte real")
    p.add_argument("--saida", type=Path, default=Path("out/consolidado.xlsx"))
    args = p.parse_args()

    if args.exemplo or args.origem is None:
        rel = dados.exemplo()
    else:
        try:
            rel = dados.carregar(args.origem)
        except NotImplementedError as erro:
            raise SystemExit(
                f"{erro}\n"
                f"Para ver o desenho a funcionar entretanto: "
                f"python relatorio/relatorio.py --exemplo"
            )

    caminho = construir(rel, args.saida)
    problemas = dados.validar(rel)
    print(f"Escrito {caminho}")
    print(f"Gráficos em {caminho.parent / 'graficos'}")
    if problemas:
        print(f"{len(problemas)} contradição(ões) — ver a folha Notas")
    else:
        print("Sem contradições entre as vistas e os dados de origem")


if __name__ == "__main__":
    main()
