"""Varredura automática: PDF -> tabelas -> Excel, sem ninguém escrever secções.

É este o caminho de `python parse.py statement.pdf`.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import parse as entrypoint
from src.excel import table_state
from src.pipeline import sweep_statement
from src.router import load_layout
from src.tables import classify, detect_tables, extract_table
from tests.make_synthetic_fixture import FIXTURE, build

LAYOUT = Path("tests/fixtures/layout.synthetic.json")


@pytest.fixture(scope="module", autouse=True)
def fixture_pdf() -> Path:
    if not FIXTURE.exists():
        build()
    return FIXTURE


@pytest.fixture(scope="module")
def sweep():
    return sweep_statement(FIXTURE, load_layout(LAYOUT))


class TestDetecao:
    def test_encontra_uma_tabela_por_titulo(self, sweep):
        titles = [result.spec.title for result in sweep.tables]
        assert titles == [
            "ACCOUNT SUMMARY",
            "PORTFOLIO HOLDINGS",
            "ACCOUNT ACTIVITY",
            "INCOME",
            "FEES AND CHARGES",
        ]

    def test_tabela_que_atravessa_a_pagina_continua_a_ser_uma(self, sweep):
        holdings = next(t for t in sweep.tables if t.spec.title == "PORTFOLIO HOLDINGS")
        assert holdings.spec.pages == [2, 3]
        assert len(holdings.data_rows) == 6

    def test_colunas_saem_do_header_da_tabela(self, sweep):
        holdings = next(t for t in sweep.tables if t.spec.title == "PORTFOLIO HOLDINGS")
        names = [c["name"] for c in holdings.spec.columns]
        assert names[0] == "description"
        assert any("market_value" in name for name in names)

    def test_totais_nao_entram_nas_linhas_de_dados(self, sweep):
        holdings = next(t for t in sweep.tables if t.spec.title == "PORTFOLIO HOLDINGS")
        assert [row["description"] for row in holdings.total_rows] == ["Total Holdings"]
        assert all("Total" not in row["description"] for row in holdings.data_rows)

    def test_parenteses_continuam_a_ser_negativos(self, sweep):
        holdings = next(t for t in sweep.tables if t.spec.title == "PORTFOLIO HOLDINGS")
        column = holdings.spec.amount_column
        short = next(r for r in holdings.data_rows if r["description"].startswith("XYZ"))
        assert short[column] == Decimal("-12500.00")

    def test_descricao_partida_e_juntada(self, sweep):
        holdings = next(t for t in sweep.tables if t.spec.title == "PORTFOLIO HOLDINGS")
        assert any("DUE 11/15/2034" in row["description"] for row in holdings.data_rows)


class TestClassificacao:
    @pytest.mark.parametrize(
        "label, expected",
        [
            ("Total", "total"),
            ("Total Holdings", "total"),
            ("NET CREDITS/(DEBITS)", "total"),
            ("Purchases", "subtotal"),
            ("Total Purchases vs Market Value", "info"),
            ("Net Value Increase/(Decrease)", "info"),
            ("Asset Class: Equities", "info"),
            ("APPLE INC COMMON STOCK", "data"),
            ("", "data"),
        ],
    )
    def test_rotulos(self, label, expected):
        assert classify(label) == expected


class TestReconciliacao:
    def test_tudo_concilia_no_documento_consistente(self, sweep):
        assert sweep.failures == []
        assert len([c for c in sweep.checks if c.passed]) == 5

    def test_estado_por_tabela(self, sweep):
        for result in sweep.tables:
            assert table_state(result.spec, sweep.checks) == "conciliado"

    def test_total_torcido_no_documento_e_apanhado(self, tmp_path):
        """Se o PDF imprimir um total que não fecha, a varredura tem de dar FAIL."""
        corrupted = build(tmp_path / "corrupto.pdf", holdings_total="1,195,590.65")
        sweep = sweep_statement(corrupted, load_layout(LAYOUT))
        holdings = next(t for t in sweep.tables if t.spec.title == "PORTFOLIO HOLDINGS")
        assert holdings.sum_of(holdings.spec.amount_column) == Decimal("1195590.55")
        assert holdings.reference_total == Decimal("1195590.65")
        assert any(c.section == "PORTFOLIO HOLDINGS" and not c.passed for c in sweep.checks)


class TestExcel:
    def test_um_comando_produz_o_livro(self, tmp_path):
        out = tmp_path / "out"
        code = entrypoint.main([str(FIXTURE), "--layout", str(LAYOUT), "--out", str(out)])
        assert code == 0

        workbook = out / f"{FIXTURE.stem}.xlsx"
        assert workbook.exists()

        book = load_workbook(workbook)
        assert book.sheetnames[:3] == ["Resumo", "Reconciliação", "Totais impressos"]
        assert len(book.sheetnames) == 3 + 5

        resumo = book["Resumo"]
        header = [cell.value for cell in resumo[1]]
        assert "soma extraída" in header and "estado" in header
        estados = {row[header.index("estado")].value for row in resumo.iter_rows(min_row=2, max_row=6)}
        assert estados == {"conciliado"}

        totais = book["Totais impressos"]
        rotulos = [row[3].value for row in totais.iter_rows(min_row=2, values_only=False)]
        assert "Total Holdings" in rotulos

    def test_saida_marcada_parcial_quando_ha_falha(self, tmp_path):
        corrupted = build(tmp_path / "corrupto.pdf", holdings_total="1,195,590.65")
        out = tmp_path / "out"
        code = entrypoint.main([str(corrupted), "--layout", str(LAYOUT), "--out", str(out)])
        assert code == 1
        assert (out / "corrupto.PARTIAL.xlsx").exists()
        assert not (out / "corrupto.xlsx").exists()

    def test_strict_nao_escreve_nada(self, tmp_path):
        corrupted = build(tmp_path / "corrupto.pdf", holdings_total="1,195,590.65")
        out = tmp_path / "out"
        code = entrypoint.main(
            [str(corrupted), "--layout", str(LAYOUT), "--out", str(out), "--strict"]
        )
        assert code == 1
        assert not out.exists()
